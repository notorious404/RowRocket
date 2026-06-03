# pdf_doc_loader.py
import os
from typing import Dict, List, Any
from PyPDF2 import PdfReader
from docx import Document

from utils import ensure_directories

def read_pdf_text(path: str) -> str:
    reader = PdfReader(path)
    parts = []
    for page in reader.pages:
        parts.append(page.extract_text() or "")
    return "\n".join(parts)

def read_docx_text(path: str) -> str:
    doc = Document(path)
    parts = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)

    for table_idx, table in enumerate(doc.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            if any(cells):
                rows.append(cells)
        if not rows:
            continue

        max_cols = max(len(row) for row in rows)
        normalized_rows = [row + [""] * (max_cols - len(row)) for row in rows]
        headers = normalized_rows[0]
        body_rows = normalized_rows[1:]

        parts.append(f"\nDOCX TABLE {table_idx}:")
        parts.append("| " + " | ".join(headers) + " |")
        parts.append("| " + " | ".join(["---"] * max_cols) + " |")
        for row in body_rows:
            parts.append("| " + " | ".join(row) + " |")

    return "\n".join(parts)


def extract_docx_tables(path: str) -> List[Dict[str, Any]]:
    doc = Document(path)
    tables = []
    file_name = os.path.basename(path)
    for table_idx, table in enumerate(doc.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            if any(cells):
                rows.append(cells)
        if not rows:
            continue

        max_cols = max(len(row) for row in rows)
        normalized_rows = [row + [""] * (max_cols - len(row)) for row in rows]
        headers = normalized_rows[0]
        body_rows = normalized_rows[1:]

        tables.append({
            "source_file": file_name,
            "title": f"DOCX Table {table_idx}",
            "index": table_idx,
            "headers": headers,
            "rows": body_rows,
        })
    return tables

def save_and_read_uploaded_files(uploaded_files) -> Dict[str, str]:
    """
    Saves uploaded files into input_pdfs/ (as requested)
    and returns {file_path: extracted_text}.
    """
    ensure_directories()
    result = {}
    for uf in uploaded_files:
        filename = uf.name
        ext = os.path.splitext(filename)[1].lower()
        save_path = os.path.join("input_pdfs", filename)

        with open(save_path, "wb") as f:
            f.write(uf.read())

        if ext == ".pdf":
            text = read_pdf_text(save_path)
        elif ext in (".docx", ".doc"):
            # basic DOCX support; for .doc you may convert externally
            text = read_docx_text(save_path)
        else:
            text = ""

        result[save_path] = text
    return result
